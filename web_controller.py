from binascii import Error as BinasciiError
from datetime import date, datetime
from io import BytesIO
import logging
import re
from zipfile import BadZipFile

from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    Request,
    UploadFile,
    status,
)
from fastapi.responses import RedirectResponse, Response, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, or_
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session

from auth_security import require_admin, require_teacher, hash_password
from auth_model import TEACHER_DESIGNATIONS, TeacherAccount
from database import get_db
from enrollment_model import (
    Subject,
    SubjectEnrollment,
    SubjectPresentation,
    TeacherEmployment,
    TeacherQualification,
    TeacherTrainingArea,
    TeacherTrainingExperience,
)
from image_utils import read_picture
from student_model import SemesterUpdateRequest, Student
from student_repository import StudentRepository
from student_service import StudentService


router = APIRouter()
templates = Jinja2Templates(directory="templates")
logger = logging.getLogger(__name__)
MAX_EXCEL_SIZE = 10 * 1024 * 1024
DEPARTMENTS = [
    "Civil Technology",
    "Computer Science and Technology",
    "Electrical Technology",
    "Food Technology",
     "RAC Technology",
    "Tourism and Hospitality Management",
]
DEPARTMENT_LOGOS = {
    "Civil Technology": "CIV",
    "Computer Science and Technology": "CST",
    "Electrical Technology": "ELC",
    "Food Technology": "FOOD",
    "RAC Technology": "RAC",
    "Tourism and Hospitality Management": "THM",
}
TEACHER_DESIGNATION_ORDER = tuple(
    designation.casefold() for designation in TEACHER_DESIGNATIONS
)
LEGACY_DEPARTMENT_NAMES = {
    "Civil": "Civil Technology",
    "Computer": "Computer Science and Technology",
    "Electrical": "Electrical Technology",
    "Food": "Food Technology",
    "RAC": "RAC Technology",
    "Tourism": "Tourism and Hospitality Management",   
}
SEMESTERS = [
    f"{number}{'st' if number == 1 else 'nd' if number == 2 else 'rd' if number == 3 else 'th'} Semester"
    for number in range(1, 9)
]
SHIFTS = ["Morning", "Day"]

OPTIONAL_EXCEL_FIELDS = {
    "phone_number",
    "father_name",
    "father_phone",
    "mother_name",
    "district",
    "upazila",
    "union_name",
    "village",
    "department",
    "semester",
    "academic_session",
    "student_id",
    "nid",
    "date_of_birth",
    "shift",
}
STUDENT_EXCEL_FIELDS = [
    ("ID", "id"),
    ("Name", "name"),
    ("Roll", "roll"),
    ("Registration", "registration"),
    ("Email", "email"),
    ("Age", "age"),
    ("Phone Number", "phone_number"),
    ("Father Name", "father_name"),
    ("Father Phone", "father_phone"),
    ("Mother Name", "mother_name"),
    ("District", "district"),
    ("Upazila", "upazila"),
    ("Union", "union_name"),
    ("Village", "village"),
    ("Department", "department"),
    ("Semester", "semester"),
    ("Academic Session", "academic_session"),
    ("Other Student ID", "student_id"),
    ("NID", "nid"),
    ("Date of Birth", "date_of_birth"),
    ("Shift", "shift"),
    ("Password", None),
]


def format_student_excel_sheet(sheet) -> None:
    fields = [label for label, _ in STUDENT_EXCEL_FIELDS]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(fields)).coordinate}"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="3451C6")
    for column_index, label in enumerate(fields, start=1):
        column = sheet.cell(row=1, column=column_index).column_letter
        sheet.column_dimensions[column].width = max(16, min(len(label) + 4, 26))
    password_column = fields.index("Password") + 1
    for row in range(2, max(sheet.max_row, 101) + 1):
        sheet.cell(row=row, column=password_column).number_format = "@"
    date_column = fields.index("Date of Birth") + 1
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=date_column).number_format = "yyyy-mm-dd"


def create_service(db: Session) -> StudentService:
    return StudentService(StudentRepository(db))


def find_students(
    db: Session,
    search: str = "",
    department: str = "",
    semester: str = "",
    academic_session: str = "",
    shift: str = "",
) -> list[Student]:
    query = db.query(Student).filter(Student.account_status == "approved")
    if search:
        term = f"%{search}%"
        query = query.filter(
            or_(
                Student.name.ilike(term),
                Student.roll.ilike(term),
                Student.registration.ilike(term),
                Student.email.ilike(term),
                Student.phone_number.ilike(term),
                Student.department.ilike(term),
                Student.semester.ilike(term),
                Student.academic_session.ilike(term),
                Student.shift.ilike(term),
                Student.district.ilike(term),
                Student.upazila.ilike(term),
                Student.village.ilike(term),
                Student.father_name.ilike(term),
                Student.father_phone.ilike(term),
                Student.student_id.ilike(term),
                Student.nid.ilike(term),
            )
        )
    if department:
        query = query.filter(Student.department == department)
    if semester:
        query = query.filter(Student.semester == semester)
    if academic_session:
        query = query.filter(Student.academic_session == academic_session)
    if shift:
        query = query.filter(Student.shift == shift)
    return query.order_by(Student.id.desc()).all()


def read_students_excel(file_data: bytes) -> tuple[list[Student], list[str]]:
    try:
        workbook = load_workbook(BytesIO(file_data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, BinasciiError, OSError) as error:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is not a valid .xlsx workbook.",
        ) from error

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        workbook.close()
        raise HTTPException(status_code=400, detail="The Excel file is empty.")

    header_map = {}
    for index, value in enumerate(headers):
        if value is None:
            continue
        key = str(value).strip().lower().replace(" ", "_")
        aliases = {
            "phone": "phone_number",
            "father_phone_number": "father_phone",
            "union": "union_name",
            "session": "academic_session",
            "dob": "date_of_birth",
            "student_roll_id": "roll",
            "login_password": "password",
            "student_password": "password",
        }
        header_map[aliases.get(key, key)] = index
    required = {"name", "roll", "registration", "email", "age"}
    missing = required - header_map.keys()
    if missing:
        workbook.close()
        raise HTTPException(
            status_code=400,
            detail=f"Missing required columns: {', '.join(sorted(missing))}",
        )

    students: list[Student] = []
    errors: list[str] = []
    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue

        values = {
            key: row[index] if index < len(row) else None
            for key, index in header_map.items()
            if key in required | OPTIONAL_EXCEL_FIELDS | {"password"}
        }
        text_values = {
            key: str(values[key]).strip() if values[key] is not None else ""
            for key in ("name", "roll", "registration", "email")
        }
        try:
            age_value = values["age"]
            age = int(age_value)
            if isinstance(age_value, bool) or age <= 0 or float(age_value) != age:
                raise ValueError
        except (TypeError, ValueError, OverflowError):
            errors.append(f"Row {row_number}: age must be a positive whole number.")
            continue

        empty_fields = [key for key, value in text_values.items() if not value]
        if empty_fields:
            errors.append(
                f"Row {row_number}: missing {', '.join(empty_fields)}."
            )
            continue
        if len(text_values["name"]) > 255 or len(text_values["roll"]) > 255:
            errors.append(
                f"Row {row_number}: name and roll must be at most 255 characters."
            )
            continue
        if len(text_values["registration"]) > 255 or len(text_values["email"]) > 255:
            errors.append(
                f"Row {row_number}: registration and email must be at most 255 characters."
            )
            continue

        optional_values = {
            key: str(values[key]).strip() if values.get(key) is not None else None
            for key in OPTIONAL_EXCEL_FIELDS - {"date_of_birth"}
            if key in values
        }
        oversized = [
            key
            for key, value in optional_values.items()
            if value is not None
            and len(value) > getattr(Student, key).property.columns[0].type.length
        ]
        if oversized:
            errors.append(
                f"Row {row_number}: values too long for {', '.join(oversized)}."
            )
            continue
        dob_value = values.get("date_of_birth")
        if dob_value:
            try:
                optional_values["date_of_birth"] = (
                    dob_value.date()
                    if isinstance(dob_value, datetime)
                    else dob_value
                    if isinstance(dob_value, date)
                    else date.fromisoformat(str(dob_value).strip())
                )
            except ValueError:
                errors.append(f"Row {row_number}: date of birth must use YYYY-MM-DD.")
                continue
        else:
            optional_values["date_of_birth"] = None

        password = values.get("password")
        if password is not None:
            password = str(password).strip()
            if password and (
                len(password) != 4
                or not password.isascii()
                or not password.isdigit()
            ):
                errors.append(
                    f"Row {row_number}: password must contain exactly 4 digits (0-9)."
                )
                continue

        student_values = {
            **text_values,
            "age": age,
            **optional_values,
        }
        if password:
            student_values["account_status"] = "pending"
            student_values["password_hash"] = hash_password(password)
        students.append(Student(**student_values))

    workbook.close()
    if not students and not errors:
        errors.append("The Excel file has no student rows.")
    return students, errors


def student_or_404(service: StudentService, student_id: int) -> Student:
    student = service.get_student_by_id(student_id)
    if student is None:
        raise HTTPException(status_code=404, detail="Student not found")
    return student


def age_from_birth_date(birth_date: date) -> int:
    today = date.today()
    return today.year - birth_date.year - (
        (today.month, today.day) < (birth_date.month, birth_date.day)
    )


def student_form(
    request: Request,
    *,
    student: Student | None,
    action: str,
    title: str,
    error: str | None = None,
    form_data: dict | None = None,
    registration: bool = False,
):
    return templates.TemplateResponse(
        request=request,
        name="student_form.html",
        context={
            "student": student,
            "action": action,
            "title": title,
            "error": error,
            "form_data": form_data or {},
            "departments": DEPARTMENTS,
            "semesters": SEMESTERS,
            "shifts": SHIFTS,
            "today": date.today().isoformat(),
            "registration": registration,
        },
    )


@router.get("/")
def home(request: Request, db: Session = Depends(get_db)):
    if request.session.get("role") == "admin":
        return RedirectResponse("/admin/dashboard", status_code=status.HTTP_303_SEE_OTHER)
    if request.session.get("role") == "student":
        return RedirectResponse("/student/dashboard", status_code=status.HTTP_303_SEE_OTHER)
    if request.session.get("role") == "teacher":
        return RedirectResponse("/teacher/dashboard", status_code=status.HTTP_303_SEE_OTHER)
    return render_home(request, db)


def department_student_counts(db: Session) -> dict[str, int]:
    stored_counts = dict(
        db.query(Student.department, func.count(Student.id))
        .filter(Student.account_status == "approved")
        .filter(Student.department.is_not(None), Student.department != "")
        .group_by(Student.department)
        .all()
    )
    counts = {department: 0 for department in DEPARTMENTS}
    for department, count in stored_counts.items():
        canonical_name = LEGACY_DEPARTMENT_NAMES.get(department, department)
        if canonical_name in counts:
            counts[canonical_name] += count
    return counts


def render_home(request: Request, db: Session):
    counts = department_student_counts(db)
    return templates.TemplateResponse(
        request=request,
        name="home.html",
        context={
            "total_students": (
                db.query(func.count(Student.id))
                .filter(Student.account_status == "approved")
                .scalar()
                or 0
            ),
            "total_teachers": (
                db.query(func.count(TeacherAccount.id))
                .filter(TeacherAccount.account_status == "approved")
                .scalar()
                or 0
            ),
            "department_counts": [
                (department, DEPARTMENT_LOGOS[department], counts.get(department, 0))
                for department in DEPARTMENTS
            ],
        },
    )


@router.get("/library")
def public_library(request: Request, db: Session = Depends(get_db)):
    subject_rows = (
        db.query(Subject, TeacherAccount)
        .join(TeacherAccount, TeacherAccount.id == Subject.teacher_id)
        .filter(
            Subject.is_active.is_(True),
            TeacherAccount.account_status == "approved",
        )
        .order_by(
            Subject.department,
            TeacherAccount.name,
            Subject.semester,
            Subject.name,
        )
        .all()
    )
    departments: dict[str, dict] = {}
    for subject, teacher in subject_rows:
        department = departments.setdefault(
            subject.department,
            {"name": subject.department, "teachers": [], "_teacher_index": {}},
        )
        teacher_entry = department["_teacher_index"].get(teacher.id)
        if teacher_entry is None:
            teacher_entry = {"teacher": teacher, "subjects": []}
            department["_teacher_index"][teacher.id] = teacher_entry
            department["teachers"].append(teacher_entry)
        teacher_entry["subjects"].append(subject)

    department_groups = []
    for department in departments.values():
        department.pop("_teacher_index")
        department_groups.append(department)
    department_groups.sort(
        key=lambda group: (
            DEPARTMENTS.index(group["name"])
            if group["name"] in DEPARTMENTS
            else len(DEPARTMENTS),
            group["name"].casefold(),
        )
    )
    return templates.TemplateResponse(
        request=request,
        name="public_library.html",
        context={"department_groups": department_groups},
    )


@router.get("/library/subjects/{subject_id}")
def public_library_subject(
    request: Request,
    subject_id: int,
    db: Session = Depends(get_db),
):
    subject_and_teacher = (
        db.query(Subject, TeacherAccount)
        .join(TeacherAccount, TeacherAccount.id == Subject.teacher_id)
        .filter(
            Subject.id == subject_id,
            Subject.is_active.is_(True),
            TeacherAccount.account_status == "approved",
        )
        .first()
    )
    if subject_and_teacher is None:
        raise HTTPException(status_code=404, detail="Subject not found")

    subject, teacher = subject_and_teacher
    presentations = (
        db.query(SubjectPresentation)
        .filter(SubjectPresentation.subject_id == subject.id)
        .order_by(
            SubjectPresentation.presentation_date.desc(),
            SubjectPresentation.title,
        )
        .all()
    )
    student_id = request.session.get("user_id")
    is_approved_student = (
        request.session.get("role") == "student"
        and db.query(Student.id)
        .filter(
            Student.id == student_id,
            Student.account_status == "approved",
            Student.password_hash.is_not(None),
        )
        .first()
        is not None
    )
    has_enrollment = (
        is_approved_student
        and db.query(SubjectEnrollment.id)
        .filter(
            SubjectEnrollment.subject_id == subject.id,
            SubjectEnrollment.student_id == student_id,
            SubjectEnrollment.status == "approved",
        )
        .first()
        is not None
    )
    return templates.TemplateResponse(
        request=request,
        name="public_library_subject.html",
        context={
            "subject": subject,
            "teacher": teacher,
            "presentations": presentations,
            "can_download": has_enrollment,
            "is_student": is_approved_student,
        },
    )


@router.get("/teachers")
def public_teacher_departments(request: Request, db: Session = Depends(get_db)):
    teacher_counts = dict(
        db.query(TeacherAccount.department, func.count(TeacherAccount.id))
        .filter(
            TeacherAccount.account_status == "approved",
            TeacherAccount.department.is_not(None),
            TeacherAccount.department != "",
        )
        .group_by(TeacherAccount.department)
        .all()
    )
    return templates.TemplateResponse(
        request=request,
        name="public_teacher_departments.html",
        context={
            "departments": [
                (
                    department,
                    DEPARTMENT_LOGOS[department],
                    teacher_counts.get(department, 0),
                )
                for department in DEPARTMENTS
            ],
        },
    )


@router.get("/teachers/{department}")
def public_department_teachers(
    request: Request,
    department: str,
    db: Session = Depends(get_db),
):
    if department not in DEPARTMENTS:
        raise HTTPException(status_code=404, detail="Department not found")
    teachers = (
        db.query(TeacherAccount)
        .filter_by(department=department, account_status="approved")
        .all()
    )
    designation_priority = {
        designation: position
        for position, designation in enumerate(TEACHER_DESIGNATION_ORDER)
    }
    teachers.sort(
        key=lambda teacher: (
            designation_priority.get((teacher.designation or "").strip().casefold(), len(designation_priority)),
            (teacher.designation or "").casefold(),
            teacher.name.casefold(),
        )
    )
    return templates.TemplateResponse(
        request=request,
        name="public_department_teachers.html",
        context={"department": department, "teachers": teachers},
    )


@router.get("/teachers/{teacher_id}/photo")
def public_teacher_photo(teacher_id: int, db: Session = Depends(get_db)):
    teacher = (
        db.query(TeacherAccount)
        .filter_by(id=teacher_id, account_status="approved")
        .first()
    )
    if teacher is None or not teacher.picture_data or not teacher.picture_type:
        raise HTTPException(status_code=404, detail="Teacher photo not found")
    return Response(
        content=teacher.picture_data,
        media_type=teacher.picture_type,
        headers={"Cache-Control": "public, max-age=3600"},
    )


@router.get(
    "/admin/teachers/{teacher_id}/photo",
    dependencies=[Depends(require_admin)],
)
def admin_teacher_photo(teacher_id: int, db: Session = Depends(get_db)):
    teacher = db.query(TeacherAccount).filter_by(id=teacher_id).first()
    if teacher is None or not teacher.picture_data or not teacher.picture_type:
        raise HTTPException(status_code=404, detail="Teacher photo not found")
    return Response(content=teacher.picture_data, media_type=teacher.picture_type)


@router.get("/admin/dashboard", dependencies=[Depends(require_admin)])
def admin_dashboard(request: Request, db: Session = Depends(get_db)):
    counts = department_student_counts(db)
    return templates.TemplateResponse(
        request=request,
        name="admin_dashboard.html",
        context={
            "total_students": (
                db.query(func.count(Student.id))
                .filter(Student.account_status == "approved")
                .scalar()
                or 0
            ),
            "active_departments": sum(count > 0 for count in counts.values()),
            "pending_count": db.query(Student).filter(
                Student.account_status == "pending"
            ).count(),
            "pending_semester_count": db.query(SemesterUpdateRequest).filter_by(
                status="pending"
            ).count(),
            "pending_teacher_count": db.query(TeacherAccount).filter_by(
                account_status="pending"
            ).count(),
            "unactivated_students": (
                db.query(Student)
                .filter(
                    Student.account_status == "approved",
                    Student.password_hash.is_(None),
                )
                .order_by(Student.id)
                .all()
            ),
            "email_error": request.query_params.get("email_error") == "1",
            "approved": request.query_params.get("approved") == "1",
            "department_counts": [
                (department, DEPARTMENT_LOGOS[department], counts[department])
                for department in DEPARTMENTS
            ],
        },
    )


@router.get("/admin/authorised-list", dependencies=[Depends(require_admin)])
def authorised_student_list(
    request: Request,
    department: str = "",
    requested_semester: str = "",
    db: Session = Depends(get_db),
):
    semester_request_query = (
        db.query(SemesterUpdateRequest, Student)
        .join(Student, Student.id == SemesterUpdateRequest.student_id)
        .filter(SemesterUpdateRequest.status == "pending")
    )
    available_request_departments = [
        value
        for (value,) in (
            db.query(Student.department)
            .join(
                SemesterUpdateRequest,
                SemesterUpdateRequest.student_id == Student.id,
            )
            .filter(
                SemesterUpdateRequest.status == "pending",
                Student.department.is_not(None),
                Student.department != "",
            )
            .distinct()
            .order_by(Student.department)
            .all()
        )
    ]
    available_requested_semesters = [
        value
        for (value,) in (
            db.query(SemesterUpdateRequest.requested_semester)
            .filter(SemesterUpdateRequest.status == "pending")
            .distinct()
            .order_by(SemesterUpdateRequest.requested_semester)
            .all()
        )
    ]
    if department not in available_request_departments:
        department = ""
    if requested_semester not in available_requested_semesters:
        requested_semester = ""
    if department:
        semester_request_query = semester_request_query.filter(
            Student.department == department
        )
    if requested_semester:
        semester_request_query = semester_request_query.filter(
            SemesterUpdateRequest.requested_semester == requested_semester
        )
    return templates.TemplateResponse(
        request=request,
        name="authorised_list.html",
        context={
            "pending_registrations": (
                db.query(Student)
                .filter(Student.account_status == "pending")
                .order_by(Student.id)
                .all()
            ),
            "pending_teachers": (
                db.query(TeacherAccount)
                .filter_by(account_status="pending")
                .order_by(TeacherAccount.id)
                .all()
            ),
            "pending_semester_requests": semester_request_query.order_by(
                SemesterUpdateRequest.created_at, SemesterUpdateRequest.id
            ).all(),
            "semester_request_department": department,
            "semester_request_semester": requested_semester,
            "semester_request_departments": available_request_departments,
            "semester_request_semesters": available_requested_semesters,
            "teacher_approved": request.query_params.get("teacher_approved") == "1",
            "teacher_approval_error": request.query_params.get("teacher_approval_error") == "1",
            "teacher_approved_count": request.query_params.get("teacher_approved_count"),
            "semester_request_approved": request.query_params.get("semester_request_approved") == "1",
            "semester_request_rejected": request.query_params.get("semester_request_rejected") == "1",
            "semester_request_error": request.query_params.get("semester_request_error") == "1",
            "semester_request_approved_count": request.query_params.get("semester_request_approved_count"),
            "email_error": request.query_params.get("email_error") == "1",
            "approved": request.query_params.get("approved") == "1",
            "manual_email_error": request.query_params.get("manual_email_error") == "1",
            "approval_error": request.query_params.get("approval_error") == "1",
            "password_error": request.query_params.get("password_error") == "1",
            "manual_approved": request.query_params.get("manual_approved") == "1",
            "excel_password_ready_count": (
                db.query(Student)
                .filter(
                    Student.account_status == "pending",
                    Student.password_hash.is_not(None),
                )
                .count()
            ),
            "password_approved_count": request.query_params.get("password_approved_count"),
            "approved_count": request.query_params.get("approved_count"),
            "failed_count": request.query_params.get("failed_count"),
        },
    )


@router.get("/student/register")
def student_registration_page(request: Request):
    return student_form(
        request,
        student=None,
        action="/student/register",
        title="New Student Registration",
        registration=True,
    )


@router.post("/student/register")
async def register_student(
    request: Request,
    name: str = Form(min_length=1, max_length=255),
    roll: str = Form(min_length=1, max_length=255),
    registration: str = Form(min_length=1, max_length=255),
    email: str = Form(min_length=1, max_length=255),
    phone_number: str = Form(min_length=1, max_length=30),
    father_name: str = Form(min_length=1, max_length=255),
    father_phone: str = Form(min_length=1, max_length=30),
    mother_name: str = Form(min_length=1, max_length=255),
    district: str = Form(min_length=1, max_length=255),
    upazila: str = Form(min_length=1, max_length=255),
    union_name: str = Form(min_length=1, max_length=255),
    village: str = Form(min_length=1, max_length=255),
    department: str = Form(min_length=1, max_length=255),
    semester: str = Form(min_length=1, max_length=50),
    academic_session: str = Form(min_length=1, max_length=100),
    student_id_value: str = Form(alias="student_id", default="", max_length=100),
    nid: str = Form(default="", max_length=50),
    date_of_birth: date = Form(),
    shift: str = Form(min_length=1, max_length=50),
    picture: UploadFile | None = File(default=None),
    db: Session = Depends(get_db),
):
    form_data = {
        "name": name, "roll": roll, "registration": registration, "email": email,
        "phone_number": phone_number, "father_name": father_name,
        "father_phone": father_phone, "mother_name": mother_name,
        "district": district, "upazila": upazila, "union_name": union_name,
        "village": village, "department": department, "semester": semester,
        "academic_session": academic_session, "student_id": student_id_value,
        "nid": nid, "date_of_birth": date_of_birth.isoformat(), "shift": shift,
    }

    def registration_form(error: str):
        return student_form(
            request,
            student=None,
            action="/student/register",
            title="New Student Registration",
            error=error,
            form_data=form_data,
            registration=True,
        )

    if department not in DEPARTMENTS or semester not in SEMESTERS or shift not in SHIFTS:
        return registration_form("Choose a valid department, semester, and shift.")
    if date_of_birth > date.today():
        return registration_form("Date of birth cannot be in the future.")
    if (
        db.query(Student)
        .filter(func.lower(Student.roll) == roll.strip().lower())
        .first()
    ):
        return registration_form("This roll number is already registered.")
    if (
        db.query(Student)
        .filter(func.lower(Student.email) == email.strip().lower())
        .first()
    ):
        return registration_form("This email address is already registered.")

    try:
        picture_data, picture_type = await read_picture(picture)
    except HTTPException as error:
        return registration_form(str(error.detail))

    student = Student(
        **{
            key: value.strip()
            for key, value in form_data.items()
            if key != "date_of_birth"
        },
        date_of_birth=date_of_birth,
        age=age_from_birth_date(date_of_birth),
        picture_data=picture_data,
        picture_type=picture_type,
        account_status="pending",
    )
    db.add(student)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return registration_form("This roll number is already registered.")
    return templates.TemplateResponse(
        request=request,
        name="registration_submitted.html",
        context={"roll": student.roll},
    )


@router.get("/ui/students", dependencies=[Depends(require_admin)])
def student_list(
    request: Request,
    search: str = "",
    department: str = "",
    semester: str = "",
    academic_session: str = "",
    shift: str = "",
    imported: int = 0,
    pending_import: int = 0,
    db: Session = Depends(get_db),
):
    department = department if department in DEPARTMENTS else ""
    semester = semester if semester in SEMESTERS else ""
    shift = shift if shift in SHIFTS else ""
    available_sessions = [
        value
        for (value,) in (
            db.query(Student.academic_session)
            .filter(
                Student.account_status == "approved",
                Student.academic_session.is_not(None),
                Student.academic_session != "",
            )
            .distinct()
            .order_by(Student.academic_session.desc())
            .all()
        )
    ]
    if academic_session not in available_sessions:
        academic_session = ""
    return templates.TemplateResponse(
        request=request,
        name="students.html",
        context={
            "students": find_students(
                db, search, department, semester, academic_session, shift
            ),
            "search": search,
            "department": department,
            "semester": semester,
            "academic_session": academic_session,
            "shift": shift,
            "departments": DEPARTMENTS,
            "semesters": SEMESTERS,
            "sessions": available_sessions,
            "shifts": SHIFTS,
            "created": request.query_params.get("created") == "1",
            "imported": imported,
            "pending_import": pending_import,
        },
    )


@router.post("/ui/import", dependencies=[Depends(require_admin)])
async def import_students(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    errors = []
    students = []
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        errors = ["Please upload an .xlsx Excel file."]
    else:
        file_data = await file.read(MAX_EXCEL_SIZE + 1)
        if len(file_data) > MAX_EXCEL_SIZE:
            errors = ["Excel file must be smaller than 10 MB."]
        else:
            try:
                students, errors = read_students_excel(file_data)
            except HTTPException as error:
                errors = [str(error.detail)]

    if errors:
        return templates.TemplateResponse(
            request=request,
            name="students.html",
            context={
                "students": find_students(db),
                "search": "",
                "import_errors": errors,
                "imported": 0,
                "pending_import": 0,
            },
            status_code=status.HTTP_400_BAD_REQUEST,
        )

    seen_rolls = set()
    service = create_service(db)
    for student in students:
        normalized_roll = student.roll.strip().lower()
        if normalized_roll in seen_rolls or service.get_student_by_roll(student.roll):
            errors.append(f"Roll number '{student.roll}' already exists.")
        seen_rolls.add(normalized_roll)
    if errors:
        return templates.TemplateResponse(
            request=request,
            name="students.html",
            context={
                "students": find_students(db),
                "search": "",
                "import_errors": errors,
                "imported": 0,
                "pending_import": 0,
            },
            status_code=status.HTTP_400_BAD_REQUEST,
        )

    db.add_all(students)
    try:
        db.commit()
    except SQLAlchemyError:
        db.rollback()
        raise

    imported_pending = any(student.account_status == "pending" for student in students)
    return RedirectResponse(
        f"/ui/students?{'pending_import' if imported_pending else 'imported'}={len(students)}",
        status_code=status.HTTP_303_SEE_OTHER,
    )


@router.get("/ui/manage", dependencies=[Depends(require_admin)])
def manage_students(
    request: Request,
    search: str = "",
    db: Session = Depends(get_db),
):
    return templates.TemplateResponse(
        request=request,
        name="manage.html",
        context={
            "students": find_students(db, search),
            "search": search,
            "password_updated": request.query_params.get("password_updated") == "1",
            "password_error": request.query_params.get("password_error") == "1",
        },
    )


@router.post(
    "/ui/manage/{student_id}/password",
    dependencies=[Depends(require_admin)],
)
def update_managed_student_password(
    student_id: int,
    password: str = Form(max_length=128),
    db: Session = Depends(get_db),
):
    if len(password) != 4 or not password.isascii() or not password.isdigit():
        return RedirectResponse(
            "/ui/manage?password_error=1",
            status_code=status.HTTP_303_SEE_OTHER,
        )

    student = (
        db.query(Student)
        .filter(Student.id == student_id, Student.account_status == "approved")
        .first()
    )
    if student is None:
        raise HTTPException(status_code=404, detail="Student not found")

    student.password_hash = hash_password(password)
    try:
        db.commit()
    except SQLAlchemyError:
        db.rollback()
        logger.exception("Could not update student password for student %s", student_id)
        return RedirectResponse(
            "/ui/manage?password_error=1",
            status_code=status.HTTP_303_SEE_OTHER,
        )
    return RedirectResponse(
        "/ui/manage?password_updated=1",
        status_code=status.HTTP_303_SEE_OTHER,
    )


@router.get("/admin/teachers", dependencies=[Depends(require_admin)])
def manage_teachers(
    request: Request,
    search: str = "",
    db: Session = Depends(get_db),
):
    query = db.query(TeacherAccount)
    if search.strip():
        term = f"%{search.strip()}%"
        query = query.filter(
            TeacherAccount.name.ilike(term)
            | TeacherAccount.username.ilike(term)
            | TeacherAccount.email.ilike(term)
            | TeacherAccount.account_status.ilike(term)
        )
    return templates.TemplateResponse(
        request=request,
        name="manage_teachers.html",
        context={
            "teachers": query.order_by(TeacherAccount.name).all(),
            "departments": DEPARTMENTS,
            "search": search,
            "updated": request.query_params.get("updated") == "1",
            "deleted": request.query_params.get("deleted") == "1",
            "status_updated": request.query_params.get("status_updated") == "1",
            "error": request.query_params.get("error"),
        },
    )


@router.get("/admin/teachers/export.xlsx", dependencies=[Depends(require_admin)])
def export_teachers(search: str = "", db: Session = Depends(get_db)):
    query = db.query(TeacherAccount)
    if search.strip():
        term = f"%{search.strip()}%"
        query = query.filter(
            TeacherAccount.name.ilike(term)
            | TeacherAccount.username.ilike(term)
            | TeacherAccount.email.ilike(term)
            | TeacherAccount.account_status.ilike(term)
        )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Teachers"
    fields = [
        ("ID", "id"),
        ("Name", "name"),
        ("Username", "username"),
        ("Email", "email"),
        ("Department", "department"),
        ("Designation", "designation"),
        ("Account Status", "account_status"),
    ]
    sheet.append([label for label, _ in fields])
    for teacher in query.order_by(TeacherAccount.name).all():
        sheet.append([getattr(teacher, attribute) for _, attribute in fields])
    for row in sheet.iter_rows(min_row=2):
        for cell in row[1:]:
            cell.data_type = "s"
    for column, width in {
        "A": 10, "B": 28, "C": 24, "D": 36, "E": 34, "F": 28, "G": 18
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{sheet.max_row}"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="teachers.xlsx"'},
    )


@router.get("/admin/teachers/print", dependencies=[Depends(require_admin)])
def print_teachers(
    request: Request,
    search: str = "",
    db: Session = Depends(get_db),
):
    query = db.query(TeacherAccount)
    if search.strip():
        term = f"%{search.strip()}%"
        query = query.filter(
            TeacherAccount.name.ilike(term)
            | TeacherAccount.username.ilike(term)
            | TeacherAccount.email.ilike(term)
            | TeacherAccount.account_status.ilike(term)
        )
    return templates.TemplateResponse(
        request=request,
        name="teacher_print.html",
        context={
            "teachers": query.order_by(TeacherAccount.name).all(),
            "search": search,
        },
    )


@router.get(
    "/admin/teachers/{teacher_id}/edit",
    dependencies=[Depends(require_admin)],
)
def edit_teacher_account(
    request: Request,
    teacher_id: int,
    db: Session = Depends(get_db),
):
    teacher = db.query(TeacherAccount).filter_by(id=teacher_id).first()
    if teacher is None:
        raise HTTPException(status_code=404, detail="Teacher not found")
    return templates.TemplateResponse(
        request=request,
        name="teacher_edit.html",
        context={
            "teacher": teacher,
            "departments": DEPARTMENTS,
            "designations": TEACHER_DESIGNATIONS,
            "qualifications": (
                db.query(TeacherQualification)
                .filter_by(teacher_id=teacher.id)
                .order_by(TeacherQualification.id)
                .all()
            ),
            "training_experiences": (
                db.query(TeacherTrainingExperience)
                .filter_by(teacher_id=teacher.id)
                .order_by(TeacherTrainingExperience.id)
                .all()
            ),
            "training_areas": (
                db.query(TeacherTrainingArea)
                .filter_by(teacher_id=teacher.id)
                .order_by(TeacherTrainingArea.id)
                .all()
            ),
            "employments": (
                db.query(TeacherEmployment)
                .filter_by(teacher_id=teacher.id)
                .order_by(TeacherEmployment.id)
                .all()
            ),
            "error": request.query_params.get("error"),
        },
    )


@router.post(
    "/admin/teachers/{teacher_id}/edit",
    dependencies=[Depends(require_admin)],
)
async def save_teacher_account(
    teacher_id: int,
    name: str = Form(min_length=1, max_length=255),
    username: str = Form(min_length=3, max_length=100),
    email: str = Form(min_length=3, max_length=255),
    department: str = Form(default="", max_length=255),
    designation: str = Form(default="", max_length=255),
    password: str = Form(default="", max_length=6),
    picture: UploadFile | None = File(default=None),
    remove_picture: str = Form(default=""),
    db: Session = Depends(get_db),
):
    teacher = db.query(TeacherAccount).filter_by(id=teacher_id).first()
    if teacher is None:
        raise HTTPException(status_code=404, detail="Teacher not found")
    if (
        not name.strip()
        or not username.strip()
        or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email.strip())
    ):
        return RedirectResponse(
            f"/admin/teachers/{teacher_id}/edit?error=invalid",
            status_code=status.HTTP_303_SEE_OTHER,
        )
    if password and (len(password) != 6 or not password.isascii() or not password.isdigit()):
        return RedirectResponse(
            f"/admin/teachers/{teacher_id}/edit?error=password",
            status_code=status.HTTP_303_SEE_OTHER,
        )
    if department and department not in DEPARTMENTS:
        return RedirectResponse(
            f"/admin/teachers/{teacher_id}/edit?error=invalid",
            status_code=status.HTTP_303_SEE_OTHER,
        )
    if designation and designation not in TEACHER_DESIGNATIONS:
        return RedirectResponse(
            f"/admin/teachers/{teacher_id}/edit?error=invalid",
            status_code=status.HTTP_303_SEE_OTHER,
        )
    try:
        picture_data, picture_type = await read_picture(picture)
    except HTTPException as exc:
        if exc.status_code in (400, 413):
            return RedirectResponse(
                f"/admin/teachers/{teacher_id}/edit?error=picture",
                status_code=status.HTTP_303_SEE_OTHER,
            )
        raise
    teacher.name = name.strip()
    teacher.username = username.strip().lower()
    teacher.email = email.strip().lower()
    teacher.department = department or None
    teacher.designation = designation.strip() or None
    if password:
        teacher.password_hash = hash_password(password)
    if picture_data is not None and picture_type is not None:
        teacher.picture_data = picture_data
        teacher.picture_type = picture_type
    elif remove_picture == "1":
        teacher.picture_data = None
        teacher.picture_type = None
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse(
            f"/admin/teachers/{teacher_id}/edit?error=duplicate",
            status_code=status.HTTP_303_SEE_OTHER,
        )
    except SQLAlchemyError:
        db.rollback()
        logger.exception("Could not update teacher account %s", teacher_id)
        return RedirectResponse(
            f"/admin/teachers/{teacher_id}/edit?error=save",
            status_code=status.HTTP_303_SEE_OTHER,
        )
    return RedirectResponse(
        "/admin/teachers?updated=1",
        status_code=status.HTTP_303_SEE_OTHER,
    )


@router.post(
    "/admin/teachers/{teacher_id}/status",
    dependencies=[Depends(require_admin)],
)
def toggle_teacher_status(teacher_id: int, db: Session = Depends(get_db)):
    teacher = db.query(TeacherAccount).filter_by(id=teacher_id).first()
    if teacher is None:
        raise HTTPException(status_code=404, detail="Teacher not found")
    if teacher.account_status == "approved":
        teacher.account_status = "inactive"
    elif teacher.account_status == "inactive":
        teacher.account_status = "approved"
    else:
        return RedirectResponse(
            "/admin/teachers?error=not-approved",
            status_code=status.HTTP_303_SEE_OTHER,
        )
    try:
        db.commit()
    except SQLAlchemyError:
        db.rollback()
        logger.exception("Could not change teacher account status for %s", teacher_id)
        return RedirectResponse(
            "/admin/teachers?error=save",
            status_code=status.HTTP_303_SEE_OTHER,
        )
    return RedirectResponse(
        "/admin/teachers?status_updated=1",
        status_code=status.HTTP_303_SEE_OTHER,
    )


@router.post(
    "/admin/teachers/{teacher_id}/delete",
    dependencies=[Depends(require_admin)],
)
def delete_teacher_account(teacher_id: int, db: Session = Depends(get_db)):
    teacher = db.query(TeacherAccount).filter_by(id=teacher_id).first()
    if teacher is None:
        raise HTTPException(status_code=404, detail="Teacher not found")
    subject_ids = db.query(Subject.id).filter_by(teacher_id=teacher_id)
    try:
        db.query(TeacherQualification).filter_by(teacher_id=teacher_id).delete(
            synchronize_session=False
        )
        db.query(TeacherTrainingExperience).filter_by(teacher_id=teacher_id).delete(
            synchronize_session=False
        )
        db.query(TeacherTrainingArea).filter_by(teacher_id=teacher_id).delete(
            synchronize_session=False
        )
        db.query(TeacherEmployment).filter_by(teacher_id=teacher_id).delete(
            synchronize_session=False
        )
        db.query(SubjectPresentation).filter_by(teacher_id=teacher_id).delete(
            synchronize_session=False
        )
        db.query(SubjectEnrollment).filter(
            SubjectEnrollment.subject_id.in_(subject_ids)
        ).delete(synchronize_session=False)
        db.query(Subject).filter_by(teacher_id=teacher_id).delete(
            synchronize_session=False
        )
        db.delete(teacher)
        db.commit()
    except SQLAlchemyError:
        db.rollback()
        logger.exception("Could not delete teacher account %s", teacher_id)
        return RedirectResponse(
            "/admin/teachers?error=delete",
            status_code=status.HTTP_303_SEE_OTHER,
        )
    return RedirectResponse(
        "/admin/teachers?deleted=1",
        status_code=status.HTTP_303_SEE_OTHER,
    )


@router.get("/ui/students/new", dependencies=[Depends(require_admin)])
def new_student(request: Request):
    return student_form(
        request,
        student=None,
        action="/ui/students/new",
        title="Add student",
    )


@router.post("/ui/students/new", dependencies=[Depends(require_admin)])
async def create_student(
    request: Request,
    name: str = Form(min_length=1, max_length=255),
    roll: str = Form(min_length=1, max_length=255),
    registration: str = Form(min_length=1, max_length=255),
    email: str = Form(min_length=1, max_length=255),
    phone_number: str = Form(min_length=1, max_length=30),
    father_name: str = Form(min_length=1, max_length=255),
    father_phone: str = Form(min_length=1, max_length=30),
    mother_name: str = Form(min_length=1, max_length=255),
    district: str = Form(min_length=1, max_length=255),
    upazila: str = Form(min_length=1, max_length=255),
    union_name: str = Form(min_length=1, max_length=255),
    village: str = Form(min_length=1, max_length=255),
    department: str = Form(min_length=1, max_length=255),
    semester: str = Form(min_length=1, max_length=50),
    academic_session: str = Form(min_length=1, max_length=100),
    student_id: str = Form(default="", max_length=100),
    nid: str = Form(default="", max_length=50),
    date_of_birth: date = Form(),
    shift: str = Form(min_length=1, max_length=50),
    student_password: str = Form(min_length=4, max_length=4),
    picture: UploadFile | None = File(default=None),
    db: Session = Depends(get_db),
):
    service = create_service(db)
    form_data = {
        "name": name, "roll": roll, "registration": registration, "email": email,
        "phone_number": phone_number, "father_name": father_name,
        "father_phone": father_phone, "mother_name": mother_name,
        "district": district, "upazila": upazila, "union_name": union_name,
        "village": village, "department": department, "semester": semester,
        "academic_session": academic_session, "student_id": student_id,
        "nid": nid, "date_of_birth": date_of_birth.isoformat(), "shift": shift,
    }
    if not student_password.isascii() or not student_password.isdigit():
        return student_form(
            request, student=None, action="/ui/students/new",
            title="Add student", error="Student password must contain exactly 4 digits (0-9).",
            form_data=form_data,
        )
    if department not in DEPARTMENTS or semester not in SEMESTERS or shift not in SHIFTS:
        return student_form(
            request, student=None, action="/ui/students/new",
            title="Add student", error="Choose a valid department, semester, and shift.",
            form_data=form_data,
        )
    if date_of_birth > date.today():
        return student_form(
            request, student=None, action="/ui/students/new",
            title="Add student", error="Date of birth cannot be in the future.",
            form_data=form_data,
        )
    if service.get_student_by_roll(roll):
        return student_form(
            request, student=None, action="/ui/students/new",
            title="Add student", error="This roll number is already in use.",
            form_data=form_data,
        )

    try:
        picture_data, picture_type = await read_picture(picture)
    except HTTPException as error:
        return student_form(
            request, student=None, action="/ui/students/new",
            title="Add student", error=str(error.detail), form_data=form_data,
        )

    student = Student(
        **{key: value.strip() for key, value in form_data.items() if key not in {"date_of_birth"}},
        date_of_birth=date_of_birth,
        age=age_from_birth_date(date_of_birth),
        picture_data=picture_data,
        picture_type=picture_type,
        password_hash=hash_password(student_password),
    )
    try:
        service.add_student(student)
    except IntegrityError:
        db.rollback()
        return student_form(
            request, student=None, action="/ui/students/new",
            title="Add student", error="This roll number is already in use.",
            form_data=form_data,
        )
    return RedirectResponse(
        "/ui/students?created=1",
        status_code=status.HTTP_303_SEE_OTHER,
    )


@router.get("/ui/students/{student_id}/edit", dependencies=[Depends(require_admin)])
def edit_student(
    request: Request,
    student_id: int,
    db: Session = Depends(get_db),
):
    student = student_or_404(create_service(db), student_id)
    return student_form(
        request,
        student=student,
        action=f"/ui/students/{student_id}/edit",
        title="Edit student",
    )


@router.post("/ui/students/{student_id}/edit", dependencies=[Depends(require_admin)])
async def save_student(
    request: Request,
    student_id: int,
    name: str = Form(min_length=1, max_length=255),
    roll: str = Form(min_length=1, max_length=255),
    registration: str = Form(min_length=1, max_length=255),
    email: str = Form(min_length=1, max_length=255),
    phone_number: str = Form(min_length=1, max_length=30),
    father_name: str = Form(min_length=1, max_length=255),
    father_phone: str = Form(min_length=1, max_length=30),
    mother_name: str = Form(min_length=1, max_length=255),
    district: str = Form(min_length=1, max_length=255),
    upazila: str = Form(min_length=1, max_length=255),
    union_name: str = Form(min_length=1, max_length=255),
    village: str = Form(min_length=1, max_length=255),
    department: str = Form(min_length=1, max_length=255),
    semester: str = Form(min_length=1, max_length=50),
    academic_session: str = Form(min_length=1, max_length=100),
    student_id_value: str = Form(alias="student_id", default="", max_length=100),
    student_password: str = Form(default="", max_length=128),
    nid: str = Form(default="", max_length=50),
    date_of_birth: date = Form(),
    shift: str = Form(min_length=1, max_length=50),
    picture: UploadFile | None = File(default=None),
    db: Session = Depends(get_db),
):
    service = create_service(db)
    existing = service.get_student_by_roll(roll)
    student = student_or_404(service, student_id)
    form_data = {
        "name": name, "roll": roll, "registration": registration, "email": email,
        "phone_number": phone_number, "father_name": father_name,
        "father_phone": father_phone, "mother_name": mother_name,
        "district": district, "upazila": upazila, "union_name": union_name,
        "village": village, "department": department, "semester": semester,
        "academic_session": academic_session, "student_id": student_id_value,
        "nid": nid, "date_of_birth": date_of_birth.isoformat(), "shift": shift,
    }
    if department not in DEPARTMENTS or semester not in SEMESTERS or shift not in SHIFTS:
        return student_form(
            request, student=student, action=f"/ui/students/{student_id}/edit",
            title="Edit student", error="Choose a valid department, semester, and shift.",
            form_data=form_data,
        )
    if date_of_birth > date.today():
        return student_form(
            request, student=student, action=f"/ui/students/{student_id}/edit",
            title="Edit student", error="Date of birth cannot be in the future.",
            form_data=form_data,
        )
    if existing and existing.id != student_id:
        return student_form(
            request, student=student, action=f"/ui/students/{student_id}/edit",
            title="Edit student", error="This roll number is already in use.",
            form_data=form_data,
        )
    try:
        picture_data, picture_type = await read_picture(picture)
    except HTTPException as error:
        return student_form(
            request, student=student, action=f"/ui/students/{student_id}/edit",
            title="Edit student", error=str(error.detail), form_data=form_data,
        )

    values = {
        key: value.strip()
        for key, value in form_data.items()
        if key != "date_of_birth"
    }
    values.update(
        date_of_birth=date_of_birth,
        age=age_from_birth_date(date_of_birth),
    )
    if picture_data is not None:
        values.update(picture_data=picture_data, picture_type=picture_type)
    if student_password:
        if not student_password.isascii() or not student_password.isdigit() or len(student_password) != 4:
            return student_form(
                request, student=student, action=f"/ui/students/{student_id}/edit",
                title="Edit student", error="Student password must contain exactly 4 digits (0-9).",
                form_data=form_data,
            )
        values["password_hash"] = hash_password(student_password)

    try:
        service.update_student(student_id, values)
    except IntegrityError:
        db.rollback()
        return student_form(
            request, student=student, action=f"/ui/students/{student_id}/edit",
            title="Edit student", error="This roll number is already in use.",
            form_data=form_data,
        )
    return RedirectResponse("/ui/manage", status_code=status.HTTP_303_SEE_OTHER)


@router.get("/ui/students/{student_id}/picture")
def student_picture(
    student_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    student = db.query(Student).filter(Student.id == student_id).first()
    if student is None or student.picture_data is None:
        raise HTTPException(status_code=404, detail="Picture not found")
    if request.session.get("role") == "admin":
        require_admin(request, db)
    elif request.session.get("role") == "student":
        viewer = db.query(Student).filter_by(id=request.session.get("user_id")).first()
        if (
            viewer is None
            or (viewer.id != student.id and (
                not viewer.department or viewer.department != student.department
            ))
        ):
            raise HTTPException(status_code=404, detail="Picture not found")
    elif request.session.get("role") == "teacher":
        require_teacher(request, db)
    else:
        raise HTTPException(status_code=401, detail="Login required")
    return Response(student.picture_data, media_type=student.picture_type)


@router.post("/ui/students/{student_id}/delete", dependencies=[Depends(require_admin)])
def remove_student(student_id: int, db: Session = Depends(get_db)):
    db.query(SubjectEnrollment).filter_by(student_id=student_id).delete(
        synchronize_session=False
    )
    if not create_service(db).delete_student(student_id):
        db.rollback()
        raise HTTPException(status_code=404, detail="Student not found")
    return RedirectResponse("/ui/manage", status_code=status.HTTP_303_SEE_OTHER)


@router.get("/ui/print", dependencies=[Depends(require_admin)])
def print_students(
    request: Request,
    search: str = "",
    department: str = "",
    semester: str = "",
    academic_session: str = "",
    shift: str = "",
    db: Session = Depends(get_db),
):
    department = department if department in DEPARTMENTS else ""
    semester = semester if semester in SEMESTERS else ""
    shift = shift if shift in SHIFTS else ""
    available_sessions = {
        value
        for (value,) in (
            db.query(Student.academic_session)
            .filter(
                Student.account_status == "approved",
                Student.academic_session.is_not(None),
                Student.academic_session != "",
            )
            .distinct()
            .all()
        )
    }
    if academic_session not in available_sessions:
        academic_session = ""
    return templates.TemplateResponse(
        request=request,
        name="print.html",
        context={
            "students": find_students(
                db, search, department, semester, academic_session, shift
            ),
            "search": search,
            "department": department,
            "semester": semester,
            "academic_session": academic_session,
            "shift": shift,
        },
    )


@router.get("/ui/import-template.xlsx", dependencies=[Depends(require_admin)])
def download_student_import_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"
    fields = [label for label, _ in STUDENT_EXCEL_FIELDS]
    sheet.append(fields)
    format_student_excel_sheet(sheet)

    instructions = workbook.create_sheet("Instructions")
    instructions.append(["Student Excel Import Template"])
    instructions.append(["Required columns: Name, Roll, Registration, Email, Age."])
    instructions.append(["Fill one student per row in the Students sheet."])
    instructions.append(["Other personal and academic columns are optional."])
    instructions.append(["Password is optional; if provided, enter exactly 4 digits (0-9)."])
    instructions.append(["Format Password cells as Text to preserve leading zeroes, for example 0123."])
    instructions.append(["Rows with a password will appear in the Authorised List for bulk approval."])
    instructions.append(["Rows without a password are imported as approved, matching the existing import behavior."])
    instructions.column_dimensions["A"].width = 100
    for cell in instructions[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="3451C6")

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="student_import_template.xlsx"'},
    )


@router.get("/ui/export.xlsx", dependencies=[Depends(require_admin)])
def export_students(search: str = "", db: Session = Depends(get_db)):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"
    fields = STUDENT_EXCEL_FIELDS
    sheet.append([label for label, _ in fields])

    for student in find_students(db, search):
        sheet.append([
            getattr(student, attribute) if attribute else None
            for _, attribute in fields
        ])
    format_student_excel_sheet(sheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="students.xlsx"'},
    )
